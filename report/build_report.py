#!/usr/bin/env python3
"""Bangun laporan UTS dalam format DOCX dan PDF.

- DOCX : pandoc (pypandoc) + penyesuaian gaya dengan python-docx
         (Times New Roman 12, hitam, rata kanan-kiri, gambar tertanam).
- PDF  : HTML + Chrome headless (print-to-pdf), gaya sama.

Jalankan dari folder report:
    ../backend/.venv/bin/python build_report.py
"""
import pathlib
import re
import subprocess
import tempfile

import markdown
import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent
MD = ROOT / "LAPORAN_UTS.md"
DOCX = ROOT / "LAPORAN_UTS.docx"
PDF = ROOT / "LAPORAN_UTS.pdf"
HTML = ROOT / "LAPORAN_UTS.html"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
BLACK = RGBColor(0, 0, 0)


# --------------------------------------------------------------------- DOCX
def force_font(run, name="Times New Roman", size=12, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def style_font(style, size=12, bold=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold


def build_docx():
    tmp = pathlib.Path(tempfile.mkdtemp()) / "raw.docx"
    pypandoc.convert_file(
        str(MD), "docx", outputfile=str(tmp),
        extra_args=[f"--resource-path={ROOT}"],
    )

    doc = Document(str(tmp))

    # Halaman A4, margin 2,5 cm.
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)

    # Gaya dasar.
    normal = doc.styles["Normal"]
    style_font(normal, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size in (("Title", 14), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        style_font(st, size, bold=True)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    try:
        style_font(doc.styles["Source Code"], 9)
    except KeyError:
        pass

    # Paragraf isi.
    for p in doc.paragraphs:
        name = p.style.name if p.style is not None else ""
        is_code = name == "Source Code"
        is_heading = name.startswith("Heading") or name == "Title"

        if is_code:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = True
            for r in p.runs:
                force_font(r, "Courier New", 9)
            continue

        if not is_heading:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for r in p.runs:
            force_font(r, "Times New Roman", 12, bold=True if is_heading else None)

    # Halaman depan: kop, judul, dan identitas.
    cover_end = next(
        (i for i, p in enumerate(doc.paragraphs)
         if p.text.strip().startswith("Nama Mahasiswa")),
        None,
    )
    if cover_end is not None:
        for idx, p in enumerate(doc.paragraphs[:cover_end]):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            size = 14 if p.text.strip().startswith("LAPORAN UJIAN") else 12
            for r in p.runs:
                force_font(r, "Times New Roman", size, bold=True)

        ident = doc.paragraphs[cover_end]
        ident.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ident.paragraph_format.line_spacing = 1.15
        ident.paragraph_format.space_before = Pt(10)
        ident.paragraph_format.space_after = Pt(10)
        for r in ident.runs:
            force_font(r, "Times New Roman", 12, bold=False)

    # Tabel: font hitam Times New Roman, header tebal.
    for table in doc.tables:
        try:
            table.style = "Table Grid"
        except KeyError:
            tbl_pr = table._tbl.tblPr
            borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = borders.makeelement(qn(f"w:{edge}"), {
                    qn("w:val"): "single", qn("w:sz"): "4",
                    qn("w:space"): "0", qn("w:color"): "000000",
                })
                borders.append(el)
            tbl_pr.append(borders)
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        force_font(r, "Times New Roman", 12, bold=True if row_idx == 0 else None)

    # Gambar: tengah, lebar 6 cm.
    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = Inches(2.4)
        shape.height = Inches(2.4 * ratio)

    for p in doc.paragraphs:
        if p._element.findall(".//" + qn("w:drawing")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(DOCX))
    print(f"DOCX: {DOCX} ({DOCX.stat().st_size} bytes)")


# ---------------------------------------------------------------------- PDF
CSS = """
@page { size: A4; margin: 2.5cm; }
body {
  font-family: "Times New Roman", Times, serif;
  font-size: 12pt; color: #000; text-align: justify; line-height: 1.5;
}
h1, h2, h3, h4 { color: #000; font-family: "Times New Roman", Times, serif; }
h1 { font-size: 14pt; text-align: center; }
h2 { font-size: 12pt; border-bottom: 1px solid #000; padding-bottom: 2px; page-break-after: avoid; }
h3 { font-size: 12pt; page-break-after: avoid; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0; font-size: 11pt; page-break-inside: avoid; }
th, td { border: 1px solid #000; padding: 4pt 5pt; vertical-align: top; text-align: left; }
th { font-weight: bold; background-color: #f1f5f9; }
pre, code { font-family: "Courier New", monospace; font-size: 8.5pt; color: #000; text-align: left; }
pre { border: 1px solid #000; padding: 6pt; white-space: pre-wrap; page-break-inside: auto; word-break: break-all; line-height: 1.2; }
img { display: block; width: 6.2cm; margin: 8pt auto 4pt auto; border: 1.5pt solid #000; page-break-inside: avoid; }
.cover p { text-align: center; margin: 2pt 0; font-weight: bold; }
.cover p:nth-of-type(5) { font-size: 14pt; }
.cover hr { margin: 8pt 0; }
blockquote { margin: 6pt 0; padding-left: 10pt; border-left: 3px solid #000; color: #000; }
hr { border: none; border-top: 1px solid #000; }
"""


def build_pdf():
    text = MD.read_text(encoding="utf-8")
    body = markdown.markdown(
        text, extensions=["tables", "fenced_code", "sane_lists"]
    )
    # Ubah path gambar relatif menjadi absolut agar terbaca Chrome.
    body = re.sub(
        r'src="(?!https?:|file:|data:)([^"]+)"',
        lambda m: f'src="file://{ROOT / m.group(1)}"',
        body,
    )
    # Bungkus kop + judul dalam blok cover agar rata tengah.
    idx = body.find("<p>Nama Mahasiswa")
    if idx > 0:
        body = '<div class="cover">' + body[:idx] + "</div>" + body[idx:]

    html = (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<style>{CSS}</style></head><body>{body}</body></html>"
    )
    HTML.write_text(html, encoding="utf-8")

    subprocess.run(
        [
            CHROME,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            "--virtual-time-budget=4000",
            f"--print-to-pdf={PDF}",
            f"file://{HTML}",
        ],
        check=True,
        capture_output=True,
    )
    print(f"PDF : {PDF} ({PDF.stat().st_size} bytes)")


if __name__ == "__main__":
    build_docx()
    build_pdf()
